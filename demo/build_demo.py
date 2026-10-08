#!/usr/bin/env python3
"""Build the Researchly demo manuscript in three formats from ONE source.

    python3 demo/build_demo.py [--out demo/dist]

Writes, from demo/manuscript.py:

    dist/researchly-demo.docx           Word (python-docx + OMML equations)
    dist/researchly-demo-overleaf.zip   Overleaf project (main.tex + sections/)
    dist/researchly-demo.md             Markdown, for the website's text box
    dist/figures/*.png                  the two figures (matplotlib)

The planted errors live in the manuscript's text, not in the renderers, so
the three versions carry the same ones (see demo/README.md). Python 3.10.
"""

from __future__ import annotations

import argparse
import copy
import io
import re
import sys
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import manuscript as ms  # noqa: E402

TOKEN = re.compile(r"\{(ref|citep|citet|citeyearpar|m):([^{}]*)\}")
FLOAT_WORD = {"fig": "Figure", "tab": "Table", "eq": "Equation"}


# --- the text in each format -------------------------------------------------

def float_number(key: str) -> int:
    if key in ms.FLOATS:
        return ms.FLOATS[key].number
    if key in ms.EQUATIONS:
        return ms.EQUATIONS[key].number
    if key in ms.MISSING_REFS:
        return ms.MISSING_REFS[key]
    raise KeyError(key)


def _cite_plain(keys: list[str], textual: bool) -> str:
    refs = [ms.REFERENCES[k] for k in keys]
    if textual:
        (r,) = refs
        return f"{r['short']} ({r['year']})"
    return "(" + "; ".join(f"{r['short']}, {r['year']}" for r in refs) + ")"


def render_plain(text: str) -> str:
    """The text as a Word or Markdown reader sees it (Unicode maths,
    written cross-references, author-year citations)."""
    text = text.replace("{total_cases}", f"{ms.TOTAL_CASES:,}")

    def sub(m):
        kind, arg = m.group(1), m.group(2)
        if kind == "ref":
            prefix, key = arg.split(":", 1)
            n = float_number(arg)
            word = FLOAT_WORD[prefix]
            return f"{word} ({n})" if prefix == "eq" else f"{word} {n}"
        if kind == "citep":
            return _cite_plain(arg.split(","), textual=False)
        if kind == "citet":
            return _cite_plain([arg], textual=True)
        if kind == "citeyearpar":
            return f"({ms.REFERENCES[arg]['year']})"
        if kind == "m":
            return arg.split("|", 1)[1]
        raise ValueError(kind)

    return TOKEN.sub(sub, text)


def tex_escape(s: str) -> str:
    return (s.replace("\\", "\\textbackslash{}").replace("&", "\\&")
             .replace("%", "\\%").replace("$", "\\$").replace("#", "\\#")
             .replace("_", "\\_").replace("'", "'"))


def render_tex(text: str) -> str:
    text = text.replace("{total_cases}", f"{ms.TOTAL_CASES:,}")
    out, pos = [], 0
    for m in TOKEN.finditer(text):
        out.append(tex_escape(text[pos:m.start()]))
        kind, arg = m.group(1), m.group(2)
        if kind == "ref":
            prefix, key = arg.split(":", 1)
            if prefix == "eq":
                out.append(f"Equation~\\eqref{{{arg}}}")
            else:
                out.append(f"{FLOAT_WORD[prefix]}~\\ref{{{arg}}}")
        elif kind == "citep":
            out.append(f"\\citep{{{arg}}}")
        elif kind == "citet":
            out.append(f"\\citet{{{arg}}}")
        elif kind == "citeyearpar":
            out.append(f"\\citeyearpar{{{arg}}}")
        elif kind == "m":
            out.append("$" + arg.split("|", 1)[0] + "$")
        pos = m.end()
    out.append(tex_escape(text[pos:]))
    return "".join(out)


def caption_plain(f) -> str:
    word = "Figure" if isinstance(f, ms.Fig) else "Table"
    return f"{word} {f.number}. {f.caption}"


def reference_plain(key: str) -> str:
    r = ms.REFERENCES[key]
    authors = r["author"].split(" and ")
    names = [a.split(", ")[0] + " " + "".join(p[0] + "." for p in
                                              a.split(", ")[1].split())
             for a in authors]
    if len(names) > 2:
        who = ", ".join(names[:-1]) + " and " + names[-1]
    else:
        who = " and ".join(names)
    return (f"{who} ({r['year']}). {r['title']}. {r['journal']}, "
            f"{r['volume']}, {r['pages'].replace('--', '–')}.")


def references_sorted() -> list[str]:
    return sorted(ms.REFERENCES, key=lambda k: (ms.REFERENCES[k]["short"],
                                                ms.REFERENCES[k]["year"]))


# --- figures ---------------------------------------------------------------------

def _smooth(values: list[int]) -> list[float]:
    """A three-point moving average: the "model median" of the demo."""
    n = len(values)
    out = []
    for i in range(n):
        lo, hi = max(0, i - 1), min(n, i + 2)
        out.append(sum(values[lo:hi]) / (hi - lo))
    return out


def draw_figures(figdir: Path) -> None:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    figdir.mkdir(parents=True, exist_ok=True)
    colours = {"Marrowick": "#1f5fa0", "Tessvale": "#d9772b",
               "Oldmere": "#4f9a4f"}
    weeks = ms.WEEKS

    fig, ax = plt.subplots(figsize=(6, 3.2), dpi=100)
    for d in ms.DISTRICTS:
        ax.plot(weeks, ms.CASES[d], marker="o", ms=3, lw=1.4,
                color=colours[d], label=d)
    ax.set_xlabel("Week of the season (week 1 = first week of April)")
    ax.set_ylabel("Reported cases")
    ax.set_xlim(1, 26)
    ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False)
    fig.tight_layout()
    fig.savefig(figdir / ms.FIG_CASES.file, dpi=100)
    plt.close(fig)

    fig, axes = plt.subplots(1, 3, figsize=(6, 2.6), dpi=100, sharex=True)
    for ax, d in zip(axes, ms.DISTRICTS):
        y = ms.CASES[d]
        med = _smooth(y)
        lo = [max(0.0, v - 2.2 * v ** 0.5 - 2) for v in med]
        hi = [v + 2.2 * v ** 0.5 + 2 for v in med]
        ax.fill_between(weeks, lo, hi, color=colours[d], alpha=0.18, lw=0)
        ax.plot(weeks, med, color=colours[d], lw=1.6)
        ax.plot(weeks, y, "o", ms=2.6, color="#333333")
        ax.set_title(d, fontsize=10)
        ax.set_xlim(1, 26)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=8)
    axes[0].set_ylabel("Reported cases", fontsize=9)
    axes[1].set_xlabel("Week of the season", fontsize=9)
    fig.tight_layout()
    fig.savefig(figdir / ms.FIG_FIT.file, dpi=100)
    plt.close(fig)


# --- Word -----------------------------------------------------------------------

M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _omml_paragraph(document, eq: ms.Eq):
    """A displayed equation as real Office Math (m:oMathPara) followed by
    its number after a tab. python-docx has no equation API, so the OMML
    is built with lxml and inserted into a paragraph."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from lxml import etree

    para = document.add_paragraph()
    # A right tab stop at the text margin for the equation number.
    sec = document.sections[0]
    width = sec.page_width - sec.left_margin - sec.right_margin
    ppr = para._p.get_or_add_pPr()
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(int(width / 635)))      # EMU -> twips
    tabs.append(tab)
    ppr.append(tabs)

    omath_para = etree.SubElement(para._p, f"{{{M_NS}}}oMathPara")
    omath = etree.SubElement(omath_para, f"{{{M_NS}}}oMath")
    run = etree.SubElement(omath, f"{{{M_NS}}}r")
    t = etree.SubElement(run, f"{{{M_NS}}}t")
    t.text = eq.unicode

    number = OxmlElement("w:r")
    number.append(OxmlElement("w:tab"))
    wt = OxmlElement("w:t")
    wt.text = f"({eq.number})"
    number.append(wt)
    para._p.append(number)
    return para


def build_docx(out: Path, figdir: Path) -> None:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    d = docx.Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)

    d.add_paragraph(ms.TITLE, style="Title")
    d.add_paragraph(", ".join(ms.AUTHORS))
    d.add_paragraph(ms.AFFILIATION)

    d.add_heading("Abstract", level=1)
    for block in ms.ABSTRACT:
        d.add_paragraph(render_plain(block.text))
    d.add_paragraph("Keywords: " + "; ".join(ms.KEYWORDS))

    for block in ms.BODY:
        if isinstance(block, ms.H):
            d.add_heading(block.text, level=block.level)
        elif isinstance(block, ms.P):
            d.add_paragraph(render_plain(block.text))
        elif isinstance(block, ms.Eq):
            _omml_paragraph(d, block)
        elif isinstance(block, ms.Fig):
            d.add_picture(str(figdir / block.file), width=Inches(5.8))
            d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            d.add_paragraph(caption_plain(block), style="Caption")
        elif isinstance(block, ms.Tab):
            d.add_paragraph(caption_plain(block), style="Caption")
            table = d.add_table(rows=1, cols=len(block.header),
                                style="Table Grid")
            for cell, text in zip(table.rows[0].cells, block.header):
                cell.text = text
                cell.paragraphs[0].runs[0].bold = True
            for row in block.rows:
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    cell.text = text
            d.add_paragraph("")

    d.add_heading("References", level=1)
    for key in references_sorted():
        d.add_paragraph(reference_plain(key))

    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))


# --- Overleaf ------------------------------------------------------------------

SECTION_FILES = {"Introduction": "introduction", "Methods": "methods",
                 "Results": "results", "Discussion": "discussion",
                 "Conclusion": "conclusion"}


def _tex_table(t: ms.Tab) -> str:
    cols = "l" + "r" * (len(t.header) - 1)
    lines = [r"\begin{table}[htbp]", r"\centering",
             f"\\caption{{{tex_escape(t.caption)}}}",
             f"\\label{{{t.key}}}",
             f"\\begin{{tabular}}{{{cols}}}", r"\toprule",
             " & ".join(tex_escape(h) for h in t.header) + r" \\",
             r"\midrule"]
    for row in t.rows:
        lines.append(" & ".join(tex_escape(c) for c in row) + r" \\")
    lines += [r"\bottomrule", r"\end{tabular}", r"\end{table}"]
    return "\n".join(lines)


def _tex_figure(f: ms.Fig) -> str:
    return "\n".join([
        r"\begin{figure}[htbp]", r"\centering",
        f"\\includegraphics[width=\\linewidth]{{figures/{f.file}}}",
        f"\\caption{{{tex_escape(f.caption)}}}",
        f"\\label{{{f.key}}}",
        r"\end{figure}"])


def _tex_equation(e: ms.Eq) -> str:
    return "\n".join([r"\begin{equation}", f"  {e.latex}",
                      f"  \\label{{{e.key}}}", r"\end{equation}"])


def tex_sections() -> dict[str, str]:
    """{file name: contents} for sections/*.tex, in manuscript order."""
    files: dict[str, list[str]] = {}
    current = None
    for block in ms.BODY:
        if isinstance(block, ms.H) and block.level == 1:
            current = SECTION_FILES[block.text]
            files[current] = [f"\\section{{{tex_escape(block.text)}}}", ""]
            continue
        lines = files[current]
        if isinstance(block, ms.H):
            lines += [f"\\subsection{{{tex_escape(block.text)}}}", ""]
        elif isinstance(block, ms.P):
            # The paragraph after an equation continues the sentence
            # ("where ..."): no blank line before it.
            lines += [render_tex(block.text), ""]
        elif isinstance(block, ms.Eq):
            if lines and lines[-1] == "":
                lines.pop()
            lines += [_tex_equation(block)]
        elif isinstance(block, ms.Fig):
            lines += [_tex_figure(block), ""]
        elif isinstance(block, ms.Tab):
            lines += [_tex_table(block), ""]
    return {name: "\n".join(lines).rstrip() + "\n"
            for name, lines in files.items()}


def tex_main() -> str:
    authors = " and ".join(ms.AUTHORS)
    abstract = "\n\n".join(render_tex(b.text) for b in ms.ABSTRACT)
    inputs = "\n".join(f"\\input{{sections/{name}}}"
                       for name in SECTION_FILES.values())
    return f"""\\documentclass[11pt,a4paper]{{article}}
\\usepackage[utf8]{{inputenc}}
\\usepackage[T1]{{fontenc}}
\\usepackage{{amsmath}}
\\usepackage{{graphicx}}
\\usepackage{{booktabs}}
\\usepackage[round]{{natbib}}
\\usepackage[hidelinks]{{hyperref}}

\\title{{{tex_escape(ms.TITLE)}}}
\\author{{{authors} \\\\ {tex_escape(ms.AFFILIATION)}}}
\\date{{}}

\\begin{{document}}
\\maketitle

\\begin{{abstract}}
{abstract}
\\end{{abstract}}

\\noindent\\textbf{{Keywords:}} {"; ".join(ms.KEYWORDS)}

{inputs}

\\bibliographystyle{{plainnat}}
\\bibliography{{references}}

\\end{{document}}
"""


def bibtex() -> str:
    entries = []
    for key in references_sorted():
        r = ms.REFERENCES[key]
        entries.append("\n".join([
            f"@article{{{key},",
            f"  author  = {{{r['author']}}},",
            f"  title   = {{{r['title']}}},",
            f"  journal = {{{r['journal']}}},",
            f"  year    = {{{r['year']}}},",
            f"  volume  = {{{r['volume']}}},",
            f"  pages   = {{{r['pages']}}}",
            "}"]))
    return "\n\n".join(entries) + "\n"


def overleaf_members(figdir: Path) -> dict[str, bytes]:
    files = {"main.tex": tex_main().encode("utf-8"),
             "references.bib": bibtex().encode("utf-8")}
    for name, text in tex_sections().items():
        files[f"sections/{name}.tex"] = text.encode("utf-8")
    for f in (ms.FIG_CASES, ms.FIG_FIT):
        files[f"figures/{f.file}"] = (figdir / f.file).read_bytes()
    return files


def build_zip(out: Path, figdir: Path) -> None:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in overleaf_members(figdir).items():
            zf.writestr(name, data)
    out.write_bytes(buf.getvalue())


# --- Markdown -------------------------------------------------------------------

def _md_table(t: ms.Tab) -> str:
    lines = ["| " + " | ".join(t.header) + " |",
             "|" + "|".join("---" for _ in t.header) + "|"]
    lines += ["| " + " | ".join(row) + " |" for row in t.rows]
    return "\n".join(lines)


def markdown() -> str:
    out = [f"# {ms.TITLE}", "", ", ".join(ms.AUTHORS), "", ms.AFFILIATION,
           "", "## Abstract", ""]
    for block in ms.ABSTRACT:
        out += [render_plain(block.text), ""]
    out += ["Keywords: " + "; ".join(ms.KEYWORDS), ""]
    for block in ms.BODY:
        if isinstance(block, ms.H):
            out += ["#" * (block.level + 1) + " " + block.text, ""]
        elif isinstance(block, ms.P):
            out += [render_plain(block.text), ""]
        elif isinstance(block, ms.Eq):
            out += [f"$$ {block.latex} \\qquad ({block.number}) $$", ""]
        elif isinstance(block, ms.Fig):
            out += [f"![{caption_plain(block)}](figures/{block.file})", "",
                    caption_plain(block), ""]
        elif isinstance(block, ms.Tab):
            out += [caption_plain(block), "", _md_table(block), ""]
    out += ["## References", ""]
    for key in references_sorted():
        out += [reference_plain(key), ""]
    return "\n".join(out)


# --- entry point -------------------------------------------------------------------

def build_all(out_dir: Path) -> dict[str, Path]:
    out_dir.mkdir(parents=True, exist_ok=True)
    figdir = out_dir / "figures"
    draw_figures(figdir)
    paths = {"docx": out_dir / "researchly-demo.docx",
             "zip": out_dir / "researchly-demo-overleaf.zip",
             "md": out_dir / "researchly-demo.md"}
    build_docx(paths["docx"], figdir)
    build_zip(paths["zip"], figdir)
    paths["md"].write_text(markdown(), encoding="utf-8")
    return paths


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=HERE / "dist")
    args = ap.parse_args(argv)
    for kind, path in build_all(args.out).items():
        print(f"{kind:5} {path} ({path.stat().st_size:,} bytes)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
