"""One JATS article -> three renderings of the same prose.

    renderings = convert.renderings(xml_bytes)   # [plain, docx, tex]

- **plain** (`.txt`): paragraphs separated by blank lines, section headings
  on their own lines. Citations and inline maths appear as the reader of
  the article sees them ("[3, 4]", "R0"); display maths is dropped.
- **docx**: a real Word file built with python-docx: built-in Heading 1-9
  styles, the title in "Title", real tables for `<table-wrap>`, figures as
  a picture paragraph followed by a "Caption" paragraph, inline maths as
  text, display maths as an OMML equation (`m:oMathPara`).
- **tex**: `\\section`/`\\subsection`/..., `\\begin{abstract}`, `\\cite{}`
  for `<xref ref-type="bibr">`, `$...$` from `<tex-math>` or MathML
  `alttext` (else the formula's text in `\\mathrm{}`), display maths in
  `equation`, and `figure`/`table` floats with `\\caption` and `\\label`.

These are REPACKAGED REAL PROSE: they test how each reader handles
structure (headings, captions, tables, citations, maths) on identical
sentences, so a rule that fires a different number of times across them
points at a reader bug. They do not reproduce what Word or LaTeX authors
actually do; native `docx` and `latex` manifest items test those quirks.

Each rendering carries `expected_words`, the prose words its reader should
leave unmasked, so word counts are comparable across formats: plain and
docx count citations and inline maths as words (they are text there); tex
does not (they are masked); docx also counts table-cell text, which the
Word reader keeps as text while LaTeX masks `tabular`. For docx and plain,
`expect_text` is exactly what the reader should return, and `markers` are
the citation and maths spans in it (for tex, in the source), so parity can
set aside findings that only arise where a format masks them.
"""

from __future__ import annotations

import io
import re
import struct
import zlib
from dataclasses import dataclass, field
from typing import Callable, Optional

MATHML_NS = "http://www.w3.org/1998/Math/MathML"
OMML_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
EQUATION_PLACEHOLDER = "[equation]"   # what ingest/docx.py reads for m:oMath


class ConvertError(ValueError):
    pass


# -- model ------------------------------------------------------------------

@dataclass
class Tok:
    kind: str                        # text | cite | math
    text: str                        # as the article's reader sees it
    style: frozenset = frozenset()   # i, b, sup, sub
    keys: tuple = ()                 # cite: reference ids
    tex: Optional[str] = None        # math: LaTeX source when known


@dataclass
class Block:
    kind: str                        # heading | para | figure | table | equation
    toks: list = field(default_factory=list)   # text, or caption for floats
    level: int = 1
    label: str = ""                  # "Figure 1"
    ident: str = ""                  # xml id -> \label
    rows: list = field(default_factory=list)   # table cells (text)
    tex: Optional[str] = None        # equation
    text: str = ""                   # equation, as text
    role: str = ""                   # "" | list-bullet | list-number | quote
    abstract: bool = False


@dataclass
class Article:
    title: str
    blocks: list


@dataclass
class Rendering:
    name: str                        # plain | docx | tex | native
    filename: str
    data: bytes
    expect_text: Optional[str] = None
    markers: list = field(default_factory=list)   # [(start, end)]
    expected_words: Optional[int] = None


# -- JATS parsing -----------------------------------------------------------

def _local(tag) -> str:
    return tag.rsplit("}", 1)[-1] if isinstance(tag, str) else ""


def _attr(el, name: str) -> str:
    for key, value in el.attrib.items():
        if _local(key) == name:
            return value
    return ""


def _child(el, name: str):
    for c in el:
        if _local(c.tag) == name:
            return c
    return None


def _find(el, name: str):
    for c in el.iter():
        if _local(c.tag) == name:
            return c
    return None


def _text(el) -> str:
    return re.sub(r"\s+", " ", "".join(el.itertext())).strip()


_STYLE = {"italic": "i", "bold": "b", "sup": "sup", "sub": "sub"}
# Block-level elements that JATS allows inside <p>.
_BLOCK_IN_P = {"fig", "fig-group", "table-wrap", "table-wrap-group",
               "disp-formula", "disp-formula-group", "list", "disp-quote",
               "boxed-text", "statement", "p", "def-list", "code",
               "preformat", "array", "chem-struct-wrap"}
_SKIP_INLINE = {"fn", "inline-graphic", "inline-supplementary-material",
                "target", "graphic", "media", "alternatives"}
_SKIP_BLOCK = {"title", "label", "sec-meta", "supplementary-material",
               "graphic", "media", "ref-list", "fn-group", "caption",
               "table-wrap-foot", "attrib", "permissions", "object-id",
               "alternatives", "code", "preformat", "def-list", "array",
               "chem-struct-wrap", "glossary", "notes"}


def parse_jats(xml: bytes) -> Article:
    try:
        import defusedxml.ElementTree as DET
        from defusedxml import DefusedXmlException
    except ImportError as exc:   # pragma: no cover - a requirement
        raise ConvertError("defusedxml is required") from exc
    try:
        # Entities and external references refused; a DOCTYPE naming the
        # JATS DTD (Europe PMC always sends one) is allowed but not fetched.
        root = DET.fromstring(xml)
    except (DET.ParseError, DefusedXmlException, ValueError) as exc:
        raise ConvertError(f"not parseable XML ({type(exc).__name__})") \
            from None
    if _local(root.tag) != "article":
        raise ConvertError(f"root element is <{_local(root.tag)}>, "
                           "not <article>")
    title_el = _find(root, "article-title")
    walker = _Walker()
    meta = _find(root, "article-meta")
    if meta is not None:
        abstracts = [a for a in meta if _local(a.tag) == "abstract"]
        plain = [a for a in abstracts if not _attr(a, "abstract-type")]
        chosen = (plain or abstracts)[:1]
        for a in chosen:
            walker.abstract = True
            walker.blocks.append(Block("heading", [Tok("text", "Abstract")],
                                       level=1, abstract=True))
            walker.walk(a, 2)
            walker.abstract = False
    body = _child(root, "body")
    if body is not None:
        walker.walk(body, 1)
    if not any(b.kind == "para" for b in walker.blocks):
        raise ConvertError("no paragraphs (no full text in this record)")
    return Article(_text(title_el) if title_el is not None else "",
                   walker.blocks)


class _Walker:
    def __init__(self):
        self.blocks: list[Block] = []
        self.abstract = False
        self.role = ""

    def add(self, block: Block) -> None:
        block.abstract = self.abstract
        if block.kind == "para" and not block.role:
            block.role = self.role
        self.blocks.append(block)

    # containers ------------------------------------------------------------

    def walk(self, el, level: int) -> None:
        for c in el:
            self.block(c, level)

    def block(self, c, level: int) -> None:
        tag = _local(c.tag)
        if tag == "sec":
            title = _child(c, "title")
            if title is not None:
                toks = normalise(self.inline_all(title))
                if _visible(toks):
                    self.add(Block("heading", toks, level=level))
            self.walk(c, level + 1)
        elif tag == "p":
            self.paragraph(c, level)
        elif tag == "fig":
            self.figure(c)
        elif tag == "table-wrap":
            self.table(c)
        elif tag in ("disp-formula", "disp-formula-group"):
            for f in ([c] if tag == "disp-formula" else
                      [x for x in c if _local(x.tag) == "disp-formula"]):
                self.equation(f)
        elif tag == "list":
            role = ("list-number" if _attr(c, "list-type") in
                    ("order", "arabic", "alpha-lower", "alpha-upper",
                     "roman-lower", "roman-upper") else "list-bullet")
            outer, self.role = self.role, role
            for item in c:
                if _local(item.tag) == "list-item":
                    self.walk(item, level)
            self.role = outer
        elif tag == "disp-quote":
            outer, self.role = self.role, "quote"
            self.walk(c, level)
            self.role = outer
        elif tag in _SKIP_BLOCK:
            return
        else:            # fig-group, table-wrap-group, boxed-text, statement…
            self.walk(c, level)

    def paragraph(self, p, level: int) -> None:
        toks: list[Tok] = []

        def flush():
            nonlocal toks
            norm = normalise(toks)
            if _visible(norm):
                self.add(Block("para", norm))
            toks = []

        def on_block(el):
            flush()
            self.block(el, level)

        self.inline(p, toks, frozenset(), on_block)
        flush()

    def inline_all(self, el) -> list[Tok]:
        toks: list[Tok] = []
        self.inline(el, toks, frozenset(), lambda _el: None)
        return toks

    # inline ----------------------------------------------------------------

    def inline(self, el, toks: list, style: frozenset,
               on_block: Callable) -> None:
        if el.text:
            toks.append(Tok("text", el.text, style))
        for c in el:
            tag = _local(c.tag)
            if tag in _BLOCK_IN_P:
                on_block(c)
            elif tag == "xref" and _attr(c, "ref-type") == "bibr":
                keys = tuple(_attr(c, "rid").split()) or ("ref",)
                toks.append(Tok("cite", _text(c), style, keys=keys))
            elif tag == "inline-formula" or (tag == "math"
                                            and MATHML_NS in c.tag):
                tok = _math_tok(c)
                if tok is not None:
                    toks.append(tok)
            elif tag in _STYLE:
                self.inline(c, toks, style | {_STYLE[tag]}, on_block)
            elif tag in _SKIP_INLINE:
                pass
            elif tag == "break":
                toks.append(Tok("text", " ", style))
            else:          # xref to figures, ext-link, named-content, sc, …
                self.inline(c, toks, style, on_block)
            if c.tail:
                toks.append(Tok("text", c.tail, style))

    # floats ----------------------------------------------------------------

    def caption(self, el) -> tuple[str, list]:
        label = _child(el, "label")
        cap = _child(el, "caption")
        toks: list[Tok] = []
        if cap is not None:
            for part in cap:
                if _local(part.tag) in ("title", "p"):
                    if toks:
                        toks.append(Tok("text", " "))
                    toks += self.inline_all(part)
        return (_text(label) if label is not None else ""), normalise(toks)

    def figure(self, el) -> None:
        label, toks = self.caption(el)
        self.add(Block("figure", toks, label=label, ident=_attr(el, "id")))

    def table(self, el) -> None:
        label, toks = self.caption(el)
        rows = []
        for tr in el.iter():
            if _local(tr.tag) == "tr":
                rows.append([_text(td) for td in tr
                             if _local(td.tag) in ("td", "th")])
        rows = [r for r in rows if r]
        self.add(Block("table", toks, label=label, ident=_attr(el, "id"),
                       rows=rows))

    def equation(self, el) -> None:
        tok = _math_tok(el)
        if tok is not None:
            lab = _child(el, "label")
            self.add(Block("equation", tex=tok.tex, text=tok.text,
                           ident=_attr(el, "id"),
                           label=_text(lab).strip() if lab is not None
                           else ""))


def _math_tok(el) -> Optional[Tok]:
    tex = None
    tm = _find(el, "tex-math")
    if tm is not None:
        tex = clean_tex("".join(tm.itertext()))
    mml = el if _local(el.tag) == "math" else _find(el, "math")
    if tex is None and mml is not None and _attr(mml, "alttext"):
        tex = clean_tex(_attr(mml, "alttext"))
    text = ""
    if mml is not None:
        text = re.sub(r"\s+", "", "".join(mml.itertext()))
    if not text and tex:
        text = re.sub(r"\\[A-Za-z]+|[{}^_\\$\s]", "", tex)
    if not text and not tex:
        return None
    return Tok("math", text or "x", tex=tex)


def _balanced(s: str) -> bool:
    depth = 0
    i = 0
    while i < len(s):
        ch = s[i]
        if ch == "\\":
            i += 2
            continue
        if ch == "{":
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth < 0:
                return False
        i += 1
    return depth == 0


def clean_tex(s: str) -> Optional[str]:
    """The formula inside PMC's `\\documentclass{minimal}...$$x$$` wrapper,
    without delimiters; None if it is unusable inside `$...$`."""
    m = re.search(r"\\begin\{document\}(.*)\\end\{document\}", s, re.S)
    if m:
        s = m.group(1)
    s = s.strip()
    for a, b in (("$$", "$$"), ("\\[", "\\]"), ("\\(", "\\)"), ("$", "$")):
        if s.startswith(a) and s.endswith(b) and len(s) >= len(a) + len(b):
            s = s[len(a):len(s) - len(b)].strip()
            break
    s = re.sub(r"\\(?:begin|end)\{(?:equation|displaymath|math)\*?\}", "",
               s).strip()
    if not s or not _balanced(s) or re.search(r"(?<!\\)\$", s):
        return None
    if re.search(r"\\(?:documentclass|usepackage|begin\{document)", s):
        return None
    return s


def _visible(toks: list) -> bool:
    return any(t.text.strip() for t in toks)


_SEP = re.compile(r"^[\s,;–—-]{0,4}$")


def normalise(toks: list) -> list:
    """Collapse whitespace, merge citation runs ("[1, 2]" -> one cite),
    absorb the brackets around a citation, and fix spacing at boundaries."""
    out: list[Tok] = []
    for t in toks:
        text = re.sub(r"\s+", " ", t.text)
        if t.kind == "text" and out and out[-1].kind == "text" \
                and out[-1].style == t.style:
            out[-1] = Tok("text", out[-1].text + text, t.style)
        elif text or t.kind != "text":
            out.append(Tok(t.kind, text, t.style, t.keys, t.tex))
    # Merge citations separated only by commas, dashes or spaces.
    merged: list[Tok] = []
    i = 0
    while i < len(out):
        t = out[i]
        if t.kind == "cite":
            while True:
                if (i + 2 < len(out) and out[i + 1].kind == "text"
                        and _SEP.match(out[i + 1].text)
                        and out[i + 2].kind == "cite"):
                    t = Tok("cite", t.text + out[i + 1].text + out[i + 2].text,
                            t.style, keys=t.keys + out[i + 2].keys)
                    i += 2
                elif i + 1 < len(out) and out[i + 1].kind == "cite":
                    t = Tok("cite", t.text + out[i + 1].text, t.style,
                            keys=t.keys + out[i + 1].keys)
                    i += 1
                else:
                    break
        merged.append(t)
        i += 1
    # "[" cite "]" -> one cite whose display text has the brackets.
    for j, t in enumerate(merged):
        if t.kind != "cite" or j == 0 or j + 1 >= len(merged):
            continue
        prev, nxt = merged[j - 1], merged[j + 1]
        if prev.kind != "text" or nxt.kind != "text":
            continue
        for a, b in (("[", "]"), ("(", ")")):
            pm = re.search(re.escape(a) + r"\s*$", prev.text)
            nm = re.match(r"\s*" + re.escape(b), nxt.text)
            if pm and nm:
                merged[j - 1] = Tok("text", prev.text[:pm.start()], prev.style)
                merged[j + 1] = Tok("text", nxt.text[nm.end():], nxt.style)
                merged[j] = Tok("cite", a + t.text.strip() + b, t.style,
                                keys=t.keys)
                # "shown[1]" keeps no space; "shown [1]" keeps one.
                break
    # Spacing: no leading/trailing blanks, no doubled blanks at boundaries.
    final: list[Tok] = []
    for t in merged:
        if not t.text and t.kind == "text":
            continue
        text = t.text
        prev_text = final[-1].text if final else ""
        if (not final or prev_text.endswith(" ")) and text.startswith(" "):
            text = text.lstrip(" ")
        if not text and t.kind == "text":
            continue
        final.append(Tok(t.kind, text, t.style, t.keys, t.tex))
    while final and final[-1].kind == "text" and final[-1].text.endswith(" "):
        stripped = final[-1].text.rstrip(" ")
        if stripped:
            final[-1] = Tok("text", stripped, final[-1].style)
            break
        final.pop()
    return final


# -- text builder -------------------------------------------------------------

class _Builder:
    """Accumulates blocks joined by `sep`, tracking marker spans."""

    def __init__(self, sep: str = "\n\n"):
        self.sep = sep
        self.parts: list[str] = []
        self.length = 0
        self.markers: list[tuple[int, int]] = []
        self.started = False

    def new_block(self) -> None:
        if self.started:
            self.parts.append(self.sep)
            self.length += len(self.sep)
        self.started = True

    def add(self, text: str, marker: bool = False) -> None:
        if marker:
            self.markers.append((self.length, self.length + len(text)))
        self.parts.append(text)
        self.length += len(text)

    def text(self) -> str:
        return "".join(self.parts)


def _words(s: str) -> int:
    return len(s.split())


def _display(toks: list) -> str:
    return "".join(t.text for t in toks)


def _caption_text(block: Block) -> tuple[str, list]:
    """(prefix, toks) for plain and Word captions: "Figure 1. " + caption."""
    prefix = ""
    if block.label:
        prefix = block.label if block.label.rstrip().endswith((".", ":")) \
            else block.label + "."
        if block.toks:
            prefix += " "
    return prefix, block.toks


def _prose_words(article: Article, *, with_markers: bool,
                 with_labels: bool, with_cells: bool) -> int:
    n = 0
    for b in article.blocks:
        if b.kind in ("para", "figure", "table"):
            toks = b.toks if with_markers else \
                [t for t in b.toks if t.kind == "text"]
            # Joining tokens with a space never merges two words, so a
            # masked citation between words still leaves both counted.
            n += _words(" ".join(t.text for t in toks))
            if b.kind != "para" and with_labels:
                n += _words(_caption_text(b)[0])
            if b.kind == "table" and with_cells:
                n += sum(_words(c) for row in b.rows for c in row)
    return n


# -- plain ------------------------------------------------------------------

def render_plain(article: Article) -> Rendering:
    out = _Builder()
    for b in article.blocks:
        if b.kind == "heading":
            out.new_block()
            out.add(_display(b.toks))
        elif b.kind == "para":
            out.new_block()
            for t in b.toks:
                out.add(t.text, marker=t.kind != "text")
        elif b.kind in ("figure", "table"):
            prefix, toks = _caption_text(b)
            if not prefix and not toks:
                continue
            out.new_block()
            out.add(prefix)
            for t in toks:
                out.add(t.text, marker=t.kind != "text")
        # display maths has no plain-text form; it is left out
    text = out.text() + "\n"
    return Rendering("plain", "article.txt", text.encode("utf-8"),
                     expect_text=text, markers=out.markers,
                     expected_words=_prose_words(article, with_markers=True,
                                                 with_labels=True,
                                                 with_cells=False))


# -- docx -------------------------------------------------------------------

def tiny_png() -> bytes:
    """A valid 1x1 white PNG, built here so no binary file is committed."""
    def chunk(kind: bytes, body: bytes) -> bytes:
        return (struct.pack(">I", len(body)) + kind + body
                + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF))
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    idat = zlib.compress(b"\x00\xff\xff\xff")
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat)
            + chunk(b"IEND", b""))


def _xml_escape(s: str) -> str:
    return (s.replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;"))


_DOCX_ROLE_STYLE = {"list-bullet": "List Bullet", "list-number": "List Number",
                    "quote": "Quote"}


def render_docx(article: Article) -> Rendering:
    import docx
    from docx.oxml import parse_xml
    from docx.shared import Inches

    d = docx.Document()
    expect: list[str] = []          # one entry per Word paragraph, in order
    markers: list[tuple[int, int]] = []

    def offset() -> int:
        return sum(len(t) + 2 for t in expect)

    def runs(p, toks, prefix: str = "") -> None:
        start = offset()
        pos = 0
        if prefix:
            p.add_run(prefix)
            pos += len(prefix)
        for t in toks:
            r = p.add_run(t.text)
            r.italic = True if "i" in t.style else None
            r.bold = True if "b" in t.style else None
            if "sup" in t.style:
                r.font.superscript = True
            elif "sub" in t.style:
                r.font.subscript = True
            if t.kind != "text":
                markers.append((start + pos, start + pos + len(t.text)))
            pos += len(t.text)
        expect.append(prefix + _display(toks))

    if article.title:
        d.add_heading(article.title, 0)
        expect.append(article.title)
    png = tiny_png()
    for b in article.blocks:
        if b.kind == "heading":
            d.add_heading(_display(b.toks), min(b.level, 9))
            expect.append(_display(b.toks))
        elif b.kind == "para":
            runs(d.add_paragraph(style=_DOCX_ROLE_STYLE.get(b.role)), b.toks)
        elif b.kind == "figure":
            d.add_paragraph().add_run().add_picture(io.BytesIO(png),
                                                    width=Inches(0.5))
            expect.append("")
            prefix, toks = _caption_text(b)
            if prefix or toks:
                runs(d.add_paragraph(style="Caption"), toks, prefix)
        elif b.kind == "table":
            prefix, toks = _caption_text(b)
            if prefix or toks:
                runs(d.add_paragraph(style="Caption"), toks, prefix)
            if b.rows:
                ncols = max(len(r) for r in b.rows)
                table = d.add_table(rows=len(b.rows), cols=ncols)
                for i, row in enumerate(b.rows):
                    for j in range(ncols):
                        cell = row[j] if j < len(row) else ""
                        table.cell(i, j).text = cell
                        expect.append(cell)
        elif b.kind == "equation":
            p = d.add_paragraph()
            p._p.append(parse_xml(
                f'<m:oMathPara xmlns:m="{OMML_NS}"><m:oMath><m:r><m:t>'
                f'{_xml_escape(b.text or "x")}</m:t></m:r></m:oMath>'
                "</m:oMathPara>"))
            expect.append(EQUATION_PLACEHOLDER)
    buf = io.BytesIO()
    d.save(buf)
    return Rendering("docx", "article.docx", buf.getvalue(),
                     expect_text="\n\n".join(expect), markers=markers,
                     expected_words=_prose_words(article, with_markers=True,
                                                 with_labels=True,
                                                 with_cells=True))


# -- LaTeX ------------------------------------------------------------------

_TEX_ESCAPE = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%",
               "$": r"\$", "#": r"\#", "_": r"\_", "{": r"\{", "}": r"\}",
               "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
_TEX_ESCAPE_RE = re.compile(r"[\\&%$#_{}~^]")
_TEX_STYLE = (("i", r"\emph{"), ("b", r"\textbf{"),
              ("sup", r"\textsuperscript{"), ("sub", r"\textsubscript{"))
_SECTIONING = {1: "section", 2: "subsection", 3: "subsubsection"}


def tex_escape(s: str) -> str:
    return _TEX_ESCAPE_RE.sub(lambda m: _TEX_ESCAPE[m.group(0)], s)


def _tex_key(k: str) -> str:
    return re.sub(r"[^A-Za-z0-9:._-]", "", k) or "ref"


def _tex_label(prefix: str, ident: str) -> str:
    return f"\\label{{{prefix}:{_tex_key(ident)}}}" if ident else ""


def render_tex(article: Article) -> Rendering:
    out = _Builder(sep="\n\n")
    out.started = True
    out.add("\\documentclass{article}\n\\usepackage{amsmath}\n"
            "\\usepackage{graphicx}\n")
    if article.title:
        out.add(f"\\title{{{tex_escape(article.title)}}}\n")
    out.add("\\begin{document}\n")
    if article.title:
        out.add("\\maketitle\n")

    def inline(toks) -> None:
        for t in toks:
            if t.kind == "cite":
                out.add("\\cite{" + ",".join(_tex_key(k) for k in t.keys)
                        + "}", marker=True)
            elif t.kind == "math":
                body = t.tex if t.tex else \
                    "\\mathrm{" + tex_escape(t.text) + "}"
                out.add("$" + body + "$", marker=True)
            else:
                text = tex_escape(t.text)
                closes = ""
                for key, opener in _TEX_STYLE:
                    if key in t.style:
                        text = opener + text
                        closes += "}"
                out.add(text + closes)

    in_abstract = False
    env = ""                       # itemize / enumerate / quote being written
    for b in article.blocks:
        want_env = {"list-bullet": "itemize", "list-number": "enumerate",
                    "quote": "quote"}.get(b.role, "") \
            if b.kind == "para" else ""
        if env and (want_env != env or (in_abstract and not b.abstract)):
            out.new_block()
            out.add(f"\\end{{{env}}}")
            env = ""
        if in_abstract and not b.abstract:
            out.new_block()
            out.add("\\end{abstract}")
            in_abstract = False
        if b.kind == "heading":
            if b.abstract and b.level == 1:
                out.new_block()
                out.add("\\begin{abstract}")
                in_abstract = True
                continue
            # Inside the abstract environment, a structured abstract's parts
            # are \paragraph headings; elsewhere the JATS depth.
            level = b.level + 2 if b.abstract else b.level
            cmd = _SECTIONING.get(level, "paragraph")
            out.new_block()
            out.add(f"\\{cmd}{{")
            inline([t for t in b.toks if t.kind == "text"]
                   or [Tok("text", _display(b.toks))])
            out.add("}")
        elif b.kind == "para":
            if want_env and not env:
                out.new_block()
                out.add(f"\\begin{{{want_env}}}")
                env = want_env
            out.new_block()
            if env in ("itemize", "enumerate"):
                out.add("\\item ")
            inline(b.toks)
        elif b.kind in ("figure", "table"):
            kind = b.kind
            out.new_block()
            out.add(f"\\begin{{{kind}}}\n\\centering\n")
            if kind == "figure":
                out.add("\\includegraphics[width=0.5\\linewidth]{figure}\n")
            if b.toks:
                out.add("\\caption{")
                inline(b.toks)
                out.add("}")
            out.add(_tex_label("fig" if kind == "figure" else "tab",
                               b.ident) + "\n")
            if kind == "table" and b.rows:
                ncols = max(len(r) for r in b.rows)
                out.add(f"\\begin{{tabular}}{{{'l' * ncols}}}\n")
                for row in b.rows:
                    cells = [tex_escape(c) for c in row]
                    cells += [""] * (ncols - len(cells))
                    out.add(" & ".join(cells) + " \\\\\n")
                out.add("\\end{tabular}\n")
            out.add(f"\\end{{{kind}}}")
        elif b.kind == "equation":
            body = b.tex if b.tex else "\\mathrm{" + tex_escape(b.text) + "}"
            out.new_block()
            start = out.length
            # Numbered only when the article numbers it (a JATS <label>):
            # an unnumbered formula rendered numbered made X301 fire on
            # every one of them (P33).
            if b.label:
                out.add("\\begin{equation}\n" + body + "\n"
                        + _tex_label("eq", b.ident) + "\\end{equation}")
            else:
                out.add("\\begin{equation*}\n" + body + "\n"
                        "\\end{equation*}")
            out.markers.append((start, out.length))
    if env:
        out.new_block()
        out.add(f"\\end{{{env}}}")
    if in_abstract:
        out.new_block()
        out.add("\\end{abstract}")
    out.new_block()
    out.add("\\end{document}\n")
    text = out.text()
    return Rendering("tex", "article.tex", text.encode("utf-8"),
                     markers=out.markers,
                     expected_words=_prose_words(article, with_markers=False,
                                                 with_labels=False,
                                                 with_cells=False))


# -- entry point ------------------------------------------------------------

def renderings(xml: bytes, formats: tuple = ("plain", "docx", "tex")
               ) -> list[Rendering]:
    article = parse_jats(xml)
    makers = {"plain": render_plain, "docx": render_docx, "tex": render_tex}
    return [makers[f](article) for f in formats]
