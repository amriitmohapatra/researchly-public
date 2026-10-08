"""Tests for the real-document bench. Synthetic text only (CLAUDE.md).

    python -m pytest ml/realdocs -q

The end-to-end test needs spaCy and en_core_web_sm and is skipped without
them; everything else needs only the engine's file readers, python-docx
and PyYAML.
"""

from __future__ import annotations

import gzip
import io
import json
import re
import sys
import tarfile
import zipfile
from pathlib import Path

import pytest

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parents[1] / "packages" / "core"))

import common  # noqa: E402
import convert  # noqa: E402
import fetch  # noqa: E402
import run  # noqa: E402

# Every sentence below was written for this test.
SYNTH_JATS = b"""<?xml version="1.0" encoding="UTF-8"?>
<!DOCTYPE article PUBLIC "-//NLM//DTD JATS (Z39.96) Journal Archiving and Interchange DTD v1.2 20190208//EN" "JATS-archivearticle1.dtd">
<article xmlns:mml="http://www.w3.org/1998/Math/MathML" xmlns:xlink="http://www.w3.org/1999/xlink" article-type="research-article">
<front><article-meta>
<title-group><article-title>A synthetic study of imaginary outbreaks</article-title></title-group>
<abstract><p>We describe a made-up outbreak in a fictional town and estimate how quickly it spread.</p></abstract>
</article-meta></front>
<body>
<sec><title>Introduction</title>
<p>Imaginary outbreaks are useful for testing software that reads articles [<xref ref-type="bibr" rid="B1">1</xref>, <xref ref-type="bibr" rid="B2">2</xref>]. The basic reproduction number <inline-formula><alternatives><tex-math id="M1">\\documentclass[12pt]{minimal}\\usepackage{amsmath}\\begin{document}$$R_0$$\\end{document}</tex-math><mml:math id="M2"><mml:msub><mml:mi>R</mml:mi><mml:mn>0</mml:mn></mml:msub></mml:math></alternatives></inline-formula> summarises how fast the fictional pathogen spreads.</p>
<p>Earlier fictional work used <italic>very</italic> small samples (<xref ref-type="bibr" rid="B3">Smith et al., 2020</xref>).</p>
</sec>
<sec><title>Methods</title>
<sec><title>Model structure</title>
<p>We fitted a compartmental model with a growth rate <inline-formula><mml:math alttext="r = 0.2"><mml:mi>r</mml:mi><mml:mo>=</mml:mo><mml:mn>0.2</mml:mn></mml:math></inline-formula> per day to the invented case counts.</p>
<disp-formula id="E1"><label>(1)</label><tex-math>\\frac{dS}{dt} = -\\beta S I</tex-math></disp-formula>
<disp-formula id="E2"><tex-math>N = S + I</tex-math></disp-formula>
<p>The model was written in a general-purpose language and checked twice.</p>
<list list-type="bullet"><list-item><p>Cases were counted each week.</p></list-item><list-item><p>Deaths were counted each month.</p></list-item></list>
</sec>
</sec>
<sec><title>Results</title>
<p>The epidemic peaked in the third week, as shown in the table below.</p>
<table-wrap id="T1"><label>Table 1</label><caption><title>Weekly counts of invented cases.</title></caption>
<table><thead><tr><th>Week</th><th>Cases</th></tr></thead><tbody><tr><td>1</td><td>12</td></tr><tr><td>2</td><td>30</td></tr></tbody></table></table-wrap>
<fig id="F1"><label>Figure 1</label><caption><p>The epidemic curve of the imaginary town.</p></caption><graphic xlink:href="f1.png"/></fig>
</sec>
<sec><title>Discussion</title>
<p>Our invented results suggest that the fictional town reacted quickly to the outbreak.</p>
</sec>
</body>
<back><ref-list><ref id="B1"><mixed-citation>Reference one.</mixed-citation></ref></ref-list></back>
</article>
"""

PHRASES = ["imaginary outbreaks are useful", "fictional pathogen spreads",
           "epidemic peaked in the third week", "town reacted quickly"]


@pytest.fixture(scope="module")
def renders():
    pytest.importorskip("docx")
    return {r.name: r for r in convert.renderings(SYNTH_JATS)}


# -- convert --------------------------------------------------------------

def test_three_renderings_carry_the_structure(renders):
    assert set(renders) == {"plain", "docx", "tex"}

    plain = renders["plain"].data.decode()
    assert "\n\nMethods\n\n" in plain and "\n\nModel structure\n\n" in plain
    assert "[1, 2]." in plain and "(Smith et al., 2020)" in plain
    assert "R0 summarises" in plain
    assert "Table 1. Weekly counts" in plain
    assert "\\frac" not in plain            # display maths is left out

    tex = renders["tex"].data.decode()
    assert "\\section{Methods}" in tex and "\\subsection{Model structure}" in tex
    assert "\\begin{abstract}" in tex and "\\end{abstract}" in tex
    assert "articles \\cite{B1,B2}." in tex
    assert "$R_0$" in tex and "$r = 0.2$" in tex
    assert "\\begin{equation}" in tex and "\\label{eq:E1}" in tex
    # unnumbered in the article, unnumbered in the rendering (P33)
    assert "\\begin{equation*}\nN = S + I" in tex
    assert "\\begin{table}" in tex and "\\caption{Weekly counts" in tex
    assert "\\label{tab:T1}" in tex and "\\label{fig:F1}" in tex
    assert "\\begin{itemize}" in tex and "\\item Cases" in tex
    assert "\\emph{very}" in tex

    import docx
    d = docx.Document(io.BytesIO(renders["docx"].data))
    styles = [p.style.name for p in d.paragraphs]
    assert "Title" in styles and "Heading 1" in styles
    assert "Heading 2" in styles and "Caption" in styles
    assert "List Bullet" in styles
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "12"
    assert b"oMathPara" in zipfile.ZipFile(
        io.BytesIO(renders["docx"].data)).read("word/document.xml")


def test_renderings_read_back_cleanly(renders):
    from researchly.ingest import load_bytes
    for name, r in renders.items():
        doc = load_bytes(r.filename, r.data)
        if r.expect_text is not None:
            assert doc.original == r.expect_text, name
        sources = run.source_texts(r.filename, r.data)
        assert run.check_invariants(doc, [], sources, r.expect_text) == {}
        words = len(doc.masked.split())
        assert abs(words - r.expected_words) <= 0.15 * r.expected_words, \
            (name, words, r.expected_words)
        sections = {h.section for h in doc.headings}
        assert {"introduction", "methods", "results",
                "discussion"} <= sections, name
        for a, b in r.markers:        # markers point at citations / maths
            assert doc.original[a:b].strip(), name


def test_marker_spans_match_citations(renders):
    plain = renders["plain"]
    text = plain.expect_text
    shown = [text[a:b] for a, b in plain.markers]
    assert "[1, 2]" in shown and "(Smith et al., 2020)" in shown
    assert "R0" in shown
    tex = renders["tex"].data.decode()
    shown = [tex[a:b] for a, b in renders["tex"].markers]
    assert "\\cite{B1,B2}" in shown and "$R_0$" in shown


def test_convert_rejects_non_articles():
    with pytest.raises(convert.ConvertError):
        convert.parse_jats(b"<html><body>Not an article.</body></html>")
    with pytest.raises(convert.ConvertError):
        convert.parse_jats(b"<article><front/></article>")   # no full text
    with pytest.raises(convert.ConvertError):
        convert.parse_jats(b"<article><body><p>Broken")


def test_clean_tex():
    assert convert.clean_tex("\\documentclass{minimal}\\begin{document}"
                             "$$x^2$$\\end{document}") == "x^2"
    assert convert.clean_tex("\\(a+b\\)") == "a+b"
    assert convert.clean_tex("{unbalanced") is None
    assert convert.tex_escape("50% & $5_x") == "50\\% \\& \\$5\\_x"


# -- invariants -----------------------------------------------------------

class _Sugg:
    def __init__(self, start, end, text, rule="X001"):
        self.start, self.end, self.text, self.rule_id = start, end, text, rule


def _latex_doc():
    from researchly.ingest import load_bytes
    src = "We estimate $x$ here and the value is small.\n"
    return src, load_bytes("a.tex", src.encode())


def test_invariants_hold_on_a_clean_document():
    src, doc = _latex_doc()
    sources = run.source_texts("a.tex", src.encode())
    start = src.index("value")
    ok = _Sugg(start, start + 5, "value")
    assert run.check_invariants(doc, [ok], sources) == {}


def test_invariant_checker_catches_a_broken_document():
    src, doc = _latex_doc()
    sources = run.source_texts("a.tex", src.encode())
    n = len(src)
    dollar = src.index("$x$")
    bad = [
        _Sugg(n - 2, n + 5, ""),                     # past the end
        _Sugg(0, 2, "Xx"),                           # text does not match
        _Sugg(dollar + 1, dollar + 8, src[dollar + 1:dollar + 8]),  # in maths
    ]
    doc.masked = doc.masked[:-1]                     # length broken
    doc.segments[0].source_start = 3                 # wrong source slice
    fails = run.check_invariants(doc, bad, sources)
    assert set(fails) == {"masked_length", "segments", "span_bounds",
                          "span_text"}
    # span_in_masked needs a sound mask, so check it on its own.
    src, doc = _latex_doc()
    fails = run.check_invariants(doc, bad[2:], sources)
    assert set(fails) == {"span_in_masked"}
    # Examples carry offsets and rule ids, never text.
    for f in fails.values():
        for ex in f["examples"]:
            assert set(ex) <= {"start", "end", "rule", "reason", "original",
                               "masked", "got_len", "want_len"}


def test_readback_mismatch_is_reported():
    src, doc = _latex_doc()
    fails = run.check_invariants(doc, [], None, expect_text=src + "extra")
    assert "readback" in fails


# -- fetch: archives ------------------------------------------------------

def _tar_gz(members: dict, symlink: str | None = None) -> bytes:
    buf = io.BytesIO()
    with tarfile.open(fileobj=buf, mode="w:gz") as tf:
        for name, body in members.items():
            info = tarfile.TarInfo(name)
            info.size = len(body)
            tf.addfile(info, io.BytesIO(body))
        if symlink:
            info = tarfile.TarInfo(symlink)
            info.type = tarfile.SYMTYPE
            info.linkname = "/etc/passwd"
            tf.addfile(info)
    return buf.getvalue()


def test_arxiv_tar_gz_becomes_an_overleaf_style_zip():
    from researchly.ingest import load_bytes
    main = (b"\\documentclass{article}\n\\begin{document}\n"
            b"\\section{Introduction}\n\\input{sections/methods}\n"
            b"\\end{document}\n")
    child = b"\\section{Methods}\nWe counted invented cases every week.\n"
    data = _tar_gz({"main.tex": main, "sections/methods.tex": child,
                    "figure.png": b"\x89PNG not really",
                    "../escape.tex": b"outside", "refs.bib": b"@misc{x}"},
                   symlink="link.tex")
    files, _notes = fetch.prepare_latex(data, "2401.00001")
    assert list(files) == ["project.zip"]
    zf = zipfile.ZipFile(io.BytesIO(files["project.zip"]))
    assert sorted(zf.namelist()) == ["main.tex", "sections/methods.tex"]
    doc = load_bytes("x.zip", files["project.zip"])
    assert "invented cases" in doc.original
    assert {s.path for s in doc.segments} == {"main.tex",
                                              "sections/methods.tex"}
    sources = run.source_texts("x.zip", files["project.zip"])
    assert run.check_invariants(doc, [], sources) == {}


def test_github_archive_strips_the_top_dir_and_reads_rnw():
    from researchly.ingest import load_bytes
    top = "paper-0123456789abcdef0123456789abcdef01234567/"
    main = (b"\\documentclass{article}\n\\begin{document}\n"
            b"\\section{Methods}\nWe simulated invented outbreaks.\n"
            b"\\end{document}\n")
    data = _tar_gz({top + "ms/ms.Rnw": main, top + "README.md": b"# x",
                    top + "ms/old.tex": b"\\section{Draft}\nOld text.\n"})
    files, notes = fetch.prepare_latex(data, "0123.tar.gz")
    zf = zipfile.ZipFile(io.BytesIO(files["project.zip"]))
    assert sorted(zf.namelist()) == ["ms/ms.tex", "ms/old.tex"]
    assert any("Sweave" in n for n in notes)
    doc = load_bytes("x.zip", files["project.zip"])
    assert doc.segments[0].path == "ms/ms.tex" == fetch.zip_name("ms/ms.Rnw")


def test_project_zip_is_repacked_like_a_tarball():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("repo-abc/paper/main.tex", b"\\documentclass{article}")
        zf.writestr("repo-abc/paper/fig.pdf", b"%PDF")
        zf.writestr("repo-abc/notes.Rnw", b"\\section{A}")
    files, notes = fetch.prepare_latex(buf.getvalue(), "x.zip")
    names = zipfile.ZipFile(io.BytesIO(files["project.zip"])).namelist()
    assert sorted(names) == ["notes.tex", "paper/main.tex"]
    with pytest.raises(fetch.FetchError):
        fetch.prepare_latex(b"PK\x03\x04 truncated", "x.zip")


def test_a_damaged_latex_zip_is_an_ordinary_refusal():
    bad = convert.Rendering("native", "project.zip", b"PK\x03\x04 truncated")
    item = dict(_item(None, fmt="latex"), kind="preprint",
                notes="main: main.tex")
    rec, _ = run.run_file(item, bad, None, None)
    assert rec["status"] == "ingest_error"
    assert rec["main_expected"] == "main.tex"


def test_two_main_candidates_are_resolved_and_compared_with_the_manifest():
    # Since P29 the zip adapter chooses (a cover letter is never the main
    # file) and tells the user; the bench checks the choice against `main:`.
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("paper.tex", b"\\documentclass{article}\n"
                    b"\\begin{document}\nA short paper.\n\\end{document}\n")
        zf.writestr("cover_letter.tex", b"\\documentclass{letter}\n"
                    b"\\begin{document}\nDear editor.\n\\end{document}\n")
    r = convert.Rendering("native", "project.zip", buf.getvalue())
    item = dict(_item(None, fmt="latex"), kind="preprint",
                notes="main: paper.tex")
    rec, _ = run.run_file(item, r, None, None)
    assert rec["status"] == "ok"
    assert rec["main_picked"] == "paper.tex" and rec["main_match"] is True


def test_safe_name_keeps_the_extension():
    name = common.safe_name("S1 " + "very long supplement title " * 8
                            + ".docx")
    assert len(name) <= 120 and name.endswith(".docx")


def test_a_list_ending_the_abstract_is_closed_inside_it():
    xml = (b"<article><front><article-meta><abstract><p>Intro.</p>"
           b"<list><list-item><p>First point.</p></list-item></list>"
           b"</abstract></article-meta></front><body><list><list-item>"
           b"<p>Body point.</p></list-item></list></body></article>")
    tex = convert.render_tex(convert.parse_jats(xml)).data.decode()
    assert tex.index("\\end{itemize}") < tex.index("\\end{abstract}")
    assert tex.count("\\begin{itemize}") == 2


def test_arxiv_single_gzipped_tex_and_pdf():
    tex = b"\\documentclass{article}\\begin{document}Hello.\\end{document}"
    files, _ = fetch.prepare_latex(gzip.compress(tex), "2401.00002")
    assert files == {"main.tex": tex}
    with pytest.raises(fetch.FetchError) as err:
        fetch.prepare_latex(b"%PDF-1.5 ...", "2401.00003")
    assert err.value.status == "no_source"


def test_pick_docx_members_from_a_zip():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("supp/Table S1.docx", b"PK fake docx")
        zf.writestr("supp/data.csv", b"a,b")
    files, _ = fetch.prepare_docx(buf.getvalue(), "supplementaryFiles",
                                  "*.docx")
    assert files == {"Table_S1.docx": b"PK fake docx"}
    with pytest.raises(fetch.FetchError):
        fetch.prepare_docx(buf.getvalue(), "x", "*.pdf")


# -- manifest and policy --------------------------------------------------

def test_example_manifest_is_valid_and_pinned():
    items = common.load_manifest(HERE / "manifest.example.yaml")
    assert len(items) >= 6
    for item in items:
        assert common.ci_permitted(item)
        assert re.search(r"/[0-9a-f]{40}/", item["url"]), item["id"]


def test_notes_and_licence_policy():
    assert common.notes_of({"notes": "pick: *.docx; PLoS Med"}) \
        == {"pick": "*.docx"}
    assert common.notes_of({"notes": "main: paper/main.tex; tarball top"}) \
        == {"main": "paper/main.tex"}
    for ok in ("CC-BY-4.0", "CC0-1.0", "Apache-2.0", "MIT", "BSD-2-Clause",
               "MPL-2.0"):
        assert common.snippets_allowed({"licence": ok}), ok
    for no in ("CC-BY-NC-4.0", "CC-BY-ND-4.0", "CC-BY-NC-ND-4.0", "unknown"):
        assert not common.snippets_allowed({"licence": no}), no


def test_manifest_rejects_an_unknown_expectation(tmp_path):
    bad = tmp_path / "m.yaml"
    bad.write_text("version: 1\nitems:\n- {id: a, kind: test-file, source: "
                   "github, format: docx, url: u, licence: MIT, "
                   "permitted_use: [ci_eval], expect: maybe}\n")
    with pytest.raises(common.ManifestError):
        common.load_manifest(bad)


# -- expectations, timeouts, the I/O watch --------------------------------

def _item(expect=None, fmt="docx"):
    return {"id": "t", "kind": "test-file", "source": "github",
            "format": fmt, "licence": "MIT", "discipline": "none",
            "expect": expect, "notes": ""}


def test_expectations_decide_what_a_refusal_means():
    broken = convert.Rendering("native", "broken.docx", b"not a zip")
    rec, out = run.run_file(_item("error"), broken, None, None)
    assert rec["status"] == "expected_error" and out is None
    rec, _ = run.run_file(_item("either"), broken, None, None)
    assert rec["status"] == "expected_error"
    for expect in (None, "parse"):
        rec, _ = run.run_file(_item(expect), broken, None, None)
        assert rec["status"] == "ingest_error"
        assert rec["error_code"] == "unreadable_file"


def test_a_crashing_reader_is_recorded_not_raised(monkeypatch):
    import researchly.ingest as ingest

    def explode(name, data):
        raise RecursionError("synthetic")
    monkeypatch.setattr(ingest, "load_bytes", explode)
    r = convert.Rendering("native", "x.docx", b"PK")
    rec, _ = run.run_file(_item("parse"), r, None, None)
    assert (rec["status"], rec["stage"], rec["error_type"]) == \
        ("crash", "read", "RecursionError")


@pytest.mark.skipif(sys.platform == "win32", reason="needs fork")
def test_a_hanging_or_dying_reader_is_a_failure_not_a_stuck_job(monkeypatch):
    import os
    import time as _time
    import researchly.ingest as ingest
    r = convert.Rendering("native", "x.docx", b"PK")

    monkeypatch.setattr(ingest, "load_bytes",
                        lambda name, data: _time.sleep(30))
    t0 = _time.monotonic()
    rec, out = run.run_isolated(_item("parse"), r, None, 0.5, 5)
    assert _time.monotonic() - t0 < 10
    assert (rec["status"], rec["stage"]) == ("timeout", "read") and out is None

    monkeypatch.setattr(ingest, "load_bytes", lambda name, data: os._exit(3))
    rec, _ = run.run_isolated(_item("parse"), r, None, 5, 5)
    assert rec["status"] == "crash" and "exit 3" in rec["error_type"]


def test_io_watch_records_file_access_outside_python(tmp_path):
    run.install_watch()
    outside = tmp_path / "secret.txt"
    outside.write_text("x")
    run._WATCH = []
    try:
        outside.read_text()
    finally:
        seen, run._WATCH = run._WATCH, None
    assert "open secret.txt" in seen
    outside.read_text()                     # not watched: nothing recorded


def test_outputs_may_not_live_in_the_repository():
    with pytest.raises(SystemExit):
        common.outside_repo(common.REPO / "ml" / "realdocs" / "cache")


# -- end to end -----------------------------------------------------------

def _spacy_ready() -> bool:
    try:
        import spacy
        spacy.load("en_core_web_sm")
        return True
    except Exception:
        return False


@pytest.mark.skipif(not _spacy_ready(), reason="needs spaCy + en_core_web_sm")
def test_end_to_end_report_holds_no_document_text(tmp_path):
    pytest.importorskip("docx")
    import yaml
    import report

    mirror = tmp_path / "mirror"
    mirror.mkdir()
    (mirror / "a.xml").write_bytes(SYNTH_JATS)
    manifest = tmp_path / "manifest.yaml"
    manifest.write_text(yaml.safe_dump({"version": 1, "items": [{
        "id": "synthetic-jats", "kind": "article", "source": "europepmc",
        "format": "jats", "url": "https://example.invalid/a.xml",
        "licence": "CC-BY-4.0", "licence_evidence": "synthetic",
        "discipline": "epidemiology", "title": "synthetic", "year": 2026,
        "permitted_use": ["ci_eval"], "notes": ""}, {
        "id": "not-for-ci", "kind": "article", "source": "europepmc",
        "format": "jats", "url": "https://example.invalid/b.xml",
        "licence": "CC-BY-NC-4.0", "permitted_use": ["private_eval"]}]}))
    cache, out = tmp_path / "cache", tmp_path / "out"
    assert fetch.main(["--manifest", str(manifest), "--cache", str(cache),
                       "--mirror", f"https://example.invalid/={mirror}"]) == 0
    assert run.main(["--manifest", str(manifest), "--cache", str(cache),
                     "--out", str(out), "--no-grammar"]) == 0
    status = report.main(["--out", str(out)])

    rep_json = (out / "report.json").read_text()
    rep_md = (out / "report.md").read_text()
    results = (out / "results.jsonl").read_text()
    rep = json.loads(rep_json)
    assert status == 0, rep["failures"]
    assert rep["totals"]["analysed"] == 3
    assert rep["totals"]["items_not_run"] == {"skipped_permitted_use": 1}
    assert set(rep["flags_per_1000"]["by_format"]) == {
        "jats/plain", "jats/docx", "jats/tex"}
    prose = [p.lower() for p in PHRASES]
    for blob in (rep_json, rep_md, results):
        low = blob.lower()
        assert not any(p in low for p in prose)
    samples = [json.loads(line) for line in
               (out / "samples.jsonl").read_text().splitlines()]
    assert samples and all(len(s["snippet"]) <= run.SNIPPET_CHARS
                           for s in samples)
